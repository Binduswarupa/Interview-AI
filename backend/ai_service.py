import json
import os
import re
import requests
import tempfile

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from dotenv import load_dotenv

load_dotenv()


def get_gemini_api_key():
    api_key = os.getenv("GEMINI_API_KEY")
    if not api_key:
        raise ValueError("GEMINI_API_KEY not found. Please add it to your .env file.")
    return api_key


def generate_interview_questions(
    job_role, skills, experience_level, interview_type, num_questions, company_context=""
):
    api_key = get_gemini_api_key()

    difficulty_note = {
        "beginner": "Mostly Easy questions, with a few Medium ones.",
        "intermediate": "A balanced mix of Easy and Medium questions.",
        "advanced": "Mostly Medium and Hard questions.",
        "senior": "Primarily Hard questions with deep architectural and leadership focus.",
    }.get(experience_level, "A balanced mix.")

    prompt = f"""You are a world-class HR consultant and senior technical interviewer.
Generate a comprehensive, professional set of interview questions for the position below.

**Job Details:**
- Job Role: {job_role}
- Required Skills: {skills if skills else "General core skills for this role"}
- Experience Level: {experience_level}
- Interview Type: {interview_type}
- Total Questions Requested: {num_questions}
{f"- Company / Context: {company_context}" if company_context else ""}

**Difficulty guidance:** {difficulty_note}

Return ONLY a valid JSON object — no markdown fences, no explanation, no extra text.
Use this exact schema:

{{
  "job_role": "{job_role}",
  "experience_level": "{experience_level}",
  "interview_type": "{interview_type}",
  "technical_questions": [
    {{
      "id": 1,
      "question": "...",
      "category": "Technical",
      "difficulty": "Easy | Medium | Hard",
      "follow_up": "...",
      "evaluation_tip": "..."
    }}
  ],
  "behavioral_questions": [
    {{
      "id": 1,
      "question": "...",
      "category": "Behavioral",
      "difficulty": "Easy | Medium | Hard",
      "follow_up": "...",
      "evaluation_tip": "..."
    }}
  ],
  "situational_questions": [
    {{
      "id": 1,
      "question": "...",
      "category": "Situational",
      "difficulty": "Easy | Medium | Hard",
      "follow_up": "...",
      "evaluation_tip": "..."
    }}
  ],
  "evaluation_criteria": {{
    "scoring_scale": "Brief description of 1–5 scoring scale",
    "key_competencies": ["competency1", "competency2", "competency3"],
    "green_flags": ["positive signal 1", "positive signal 2", "positive signal 3"],
    "red_flags": ["warning sign 1", "warning sign 2", "warning sign 3"]
  }},
  "interview_tips": ["tip1", "tip2", "tip3"]
}}

Distribution rules based on interview_type:
- "technical"  → ~70% technical, ~20% behavioral, ~10% situational
- "behavioral" → ~10% technical, ~60% behavioral, ~30% situational
- "mixed"      → ~40% technical, ~35% behavioral, ~25% situational

Total questions across all three arrays must sum to approximately {num_questions}.
Make every question highly specific to "{job_role}" and the listed skills.
Ensure follow_up and evaluation_tip are insightful and actionable."""

    url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent?key={api_key}"
    headers = {
        "Content-Type": "application/json"
    }
    payload = {
        "contents": [
            {
                "parts": [
                    {
                        "text": prompt
                    }
                ]
            }
        ]
    }

    response = requests.post(url, headers=headers, json=payload)
    if response.status_code != 200:
        raise Exception(f"Gemini API returned error {response.status_code}: {response.text}")

    res_json = response.json()
    try:
        raw = res_json["candidates"][0]["content"]["parts"][0]["text"].strip()
    except (KeyError, IndexError, TypeError) as e:
        raise Exception(f"Failed to parse Gemini API response: {response.text}")

    # Strip markdown code fences if Gemini adds them
    raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.MULTILINE)
    raw = re.sub(r"\s*```$", "", raw, flags=re.MULTILINE)
    raw = raw.strip()

    return json.loads(raw)


def generate_docx(questions_data: dict, job_role: str) -> str:
    doc = Document()

    # ── Title ──────────────────────────────────────────────────────────────
    title = doc.add_heading(f"Interview Questions: {job_role}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph(
        f"Experience Level: {questions_data.get('experience_level', 'N/A').title()}   |   "
        f"Interview Type: {questions_data.get('interview_type', 'N/A').title()}"
    )
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    def add_questions(title_text, questions):
        if not questions:
            return
        doc.add_heading(title_text, level=1)
        for i, q in enumerate(questions, 1):
            doc.add_heading(f"Q{i}. {q.get('question', '')}", level=2)
            info = doc.add_paragraph()
            info.add_run(f"Difficulty: {q.get('difficulty', 'N/A')}").bold = True
            info.add_run(f"   |   Category: {q.get('category', 'N/A')}")

            if q.get("follow_up"):
                p = doc.add_paragraph()
                p.add_run("Follow-up: ").bold = True
                p.add_run(q["follow_up"])

            if q.get("evaluation_tip"):
                p = doc.add_paragraph()
                p.add_run("Evaluation Tip: ").bold = True
                p.add_run(q["evaluation_tip"])

            doc.add_paragraph("")

    add_questions("Technical Questions", questions_data.get("technical_questions", []))
    add_questions("Behavioral Questions", questions_data.get("behavioral_questions", []))
    add_questions("Situational Questions", questions_data.get("situational_questions", []))

    # ── Evaluation Criteria ────────────────────────────────────────────────
    ec = questions_data.get("evaluation_criteria", {})
    if ec:
        doc.add_heading("Evaluation Criteria", level=1)

        if ec.get("scoring_scale"):
            p = doc.add_paragraph()
            p.add_run("Scoring Scale: ").bold = True
            p.add_run(ec["scoring_scale"])

        def add_list(heading, items, bullet="•"):
            if items:
                doc.add_heading(heading, level=2)
                for item in items:
                    doc.add_paragraph(f"{bullet} {item}")

        add_list("Key Competencies", ec.get("key_competencies", []))
        add_list("Green Flags ✓", ec.get("green_flags", []), "✓")
        add_list("Red Flags ✗", ec.get("red_flags", []), "✗")

    # ── Interview Tips ─────────────────────────────────────────────────────
    tips = questions_data.get("interview_tips", [])
    if tips:
        doc.add_heading("Interview Tips", level=1)
        for tip in tips:
            doc.add_paragraph(f"• {tip}")

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(tmp.name)
    return tmp.name
